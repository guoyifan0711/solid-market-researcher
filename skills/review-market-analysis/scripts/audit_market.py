"""Read-only evidence and release checks. Python 3.12 + bundled openpyxl.

Writes only requested audit outputs; never edits workbooks or publishes content.
"""
from __future__ import annotations

import argparse
import ast
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
import difflib
import hashlib
import html
import json
import math
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest().lower()


def text_digest(text):
    return hashlib.sha256(normalize(text).encode("utf-8")).hexdigest()


def normalize(text):
    text = html.unescape(text)
    text = re.sub(r"!\[[^]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^]]*)\]\(([^)]*)\)", r"\1 \2", text)
    text = re.sub(r"^[ \t]*(?:#{1,6}\s*|[-*]\s+|\d+\.\s+)", "", text, flags=re.M)
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = text.replace("**", "").replace("`", "").replace("__", "")
    return re.sub(r"\s+", "", text)


def document(path):
    raw = Path(path).read_text(encoding="utf-8-sig")
    images = []
    if Path(path).suffix.lower() == ".xml":
        root = ET.fromstring("<audit-root>" + raw + "</audit-root>")
        def render(e):
            if e.tag == "img":
                images.append(dict(e.attrib))
                return ""
            out = (e.text or "") + "".join(render(c) + (c.tail or "") for c in e)
            if re.fullmatch(r"h[1-6]", e.tag):
                return "\n" + "#" * int(e.tag[1]) + " " + out + "\n"
            if e.tag in {"p", "li", "tr", "table", "ul", "ol", "hr", "br"}:
                return "\n" + out + "\n"
            if e.tag in {"td", "th"}:
                return " " + out + " "
            return out
        raw = render(root)
    else:
        images = [{"src": m[2], "name": m[1]} for m in re.finditer(r"!\[([^]]*)\]\(([^)]*)\)", raw)]
    pattern = re.compile(r"^#{1,6}\s*(Page\s+(\d+)\s*/\s*(\d+)[^\n]*)", re.M)
    matches = list(pattern.finditer(raw))
    if matches and any(int(m[3]) != len(matches) for m in matches):
        raise ValueError("Page N/Total denominator does not equal parsed page count")
    if matches and [int(m[2]) for m in matches] != list(range(1, len(matches) + 1)):
        raise ValueError("Page numerators must be consecutive from 1 in document order")
    pages, order = {}, []
    for i, m in enumerate(matches):
        key = "P" + m[2]
        order.append(key)
        if key in pages:
            raise ValueError("duplicate page: " + key)
        end = matches[i + 1].start() if i + 1 < len(matches) else len(raw)
        body = raw[m.end():end]
        bc = re.search(r"^#{1,6}\s*Big conclusion\s*\n(.*?)(?=\n#{1,6}\s|\Z)", body, re.M | re.S)
        pages[key] = {"heading": m[1], "body": body, "conclusion": bc[1].strip() if bc else ""}
    return {"text": raw, "intro": raw[:matches[0].start()] if matches else raw,
            "pages": pages, "order": order, "images": images}


def arithmetic(expression, values):
    if len(expression) > 500:
        raise ValueError("expression too long")
    nodes = ast.parse(expression, mode="eval")
    if sum(1 for _ in ast.walk(nodes)) > 80:
        raise ValueError("expression too complex")
    def calc(node):
        if isinstance(node, ast.Expression):
            result = calc(node.body)
        elif isinstance(node, ast.Constant) and type(node.value) in (int, float):
            result = node.value
        elif isinstance(node, ast.Name):
            result = values[node.id]
        elif isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
            result = (-1 if isinstance(node.op, ast.USub) else 1) * calc(node.operand)
        elif isinstance(node, ast.BinOp):
            a, b = calc(node.left), calc(node.right)
            if isinstance(node.op, ast.Add): result = a + b
            elif isinstance(node.op, ast.Sub): result = a - b
            elif isinstance(node.op, ast.Mult): result = a * b
            elif isinstance(node.op, ast.Div): result = a / b
            elif isinstance(node.op, ast.Pow) and abs(b) <= 20: result = a ** b
            else: raise ValueError("unsupported operator")
        else:
            raise ValueError("unsupported expression")
        if type(result) not in (float, int) or not math.isfinite(result) or abs(result) > 1e100:
            raise ValueError("nonfinite or excessive result")
        return result
    return calc(nodes)


class Audit:
    def __init__(self, manifest_path, pages=None):
        self.manifest_path = Path(manifest_path).resolve()
        self.m = json.loads(self.manifest_path.read_text(encoding="utf-8-sig"))
        if self.m.get("schema_version") != 1:
            raise ValueError("unsupported manifest schema")
        if self.manifest_path.parent.name != "05_校验记录" or self.m.get("project_root") != "..":
            raise ValueError("schema v1 manifest must be in project/05_校验记录 with project_root '..'")
        self.root = (self.manifest_path.parent / self.m["project_root"]).resolve()
        self.rows, self.books = [], {}
        baseline = self.m["baseline"].get("path")
        self.base = document(self.path(baseline)) if baseline else {"pages": {}, "order": [], "images": [], "intro": "", "text": ""}
        self.scope = pages or [p["id"] for p in self.m["pages"]]
        if not set(self.scope).issubset({p["id"] for p in self.m["pages"]}):
            raise ValueError("unknown page in --pages")
        self.pages = [p for p in self.m["pages"] if p["id"] in self.scope]

    def all_scope(self):
        return set(self.scope) == {p["id"] for p in self.m["pages"]}

    def image_review(self, candidate):
        review = self.m.get("publication", {}).get("image_review", {})
        evidence = review.get("evidence", {})
        if not (review.get("status") == "verified" and review.get("candidate_sha256") == digest(candidate)
                and review.get("baseline_sha256") == self.m["baseline"].get("sha256")
                and set(review.get("pages", [])) == set(self.scope) and evidence.get("sha256")
                and review.get("basis")):
            return False
        return self.file(evidence, "image_review_evidence")

    def path(self, relative):
        if not isinstance(relative, str) or not relative:
            raise ValueError("artifact path not registered")
        p = (self.root / relative).resolve()
        if not p.is_relative_to(self.root):
            raise ValueError("artifact outside project: " + str(relative))
        return p

    def check(self, code, ok, detail, page="", unknown=False):
        self.rows.append({"check": code, "page": page,
                          "status": "PASS" if ok else ("NOT_CHECKED" if unknown else "FAIL"),
                          "detail": detail})
        return ok

    def file(self, record, code, page=""):
        if isinstance(record, str): record = {"path": record}
        try:
            p = self.path(record["path"])
            ok = p.is_file()
            if record.get("sha256"):
                ok = ok and digest(p) == record["sha256"].lower()
            self.check(code, ok, record["path"], page)
            return ok
        except (OSError, ValueError, KeyError) as e:
            self.check(code, False, str(e), page)
            return False

    def inventory(self):
        self.file(self.m["baseline"], "baseline_hash")
        self.check("baseline_pin", bool(self.m["baseline"].get("sha256")), "批准正文必须固定哈希")
        ids = [p["id"] for p in self.m["pages"]]
        self.check("page_map", ids == self.base["order"] and len(ids) == len(set(ids)), "清单页序与正文逐页对应")
        for page in self.pages:
            pid = page["id"]
            self.check("page_heading", normalize(page["heading"]) == normalize(self.base["pages"].get(pid, {}).get("heading", "")), page["heading"], pid)
            for item in page.get("excel", []): self.file(item, "excel_file", pid)
            for item in page.get("visuals", {}).get("files", []): self.file(item, "visual_file", pid)
            for item in page.get("audits", []): self.file(item, "audit_file", pid)
            self.check("confirmation_record", bool(page.get("confirmation", {}).get("basis")), "确认状态须有依据；此检查不将待确认变成已确认", pid)

    def cell(self, spec):
        import openpyxl
        p = self.path(spec["path"])
        k = str(p)
        if k not in self.books:
            self.books[k] = (openpyxl.load_workbook(p, read_only=True, data_only=False),
                             openpyxl.load_workbook(p, read_only=True, data_only=True))
        f, v = self.books[k]
        sheet, cell = spec["sheet"], spec["cell"]
        if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", cell): raise ValueError("single cell required")
        raw, cached = f[sheet][cell].value, v[sheet][cell].value
        if type(cached) not in (int, float) or not math.isfinite(cached):
            raise ValueError("missing/non-numeric value or formula cache: " + sheet + "!" + cell)
        if not math.isclose(cached, spec["expected_value"], rel_tol=1e-9, abs_tol=1e-10):
            raise ValueError("value mismatch: " + sheet + "!" + cell + "=" + str(cached))
        if "expected_formula" in spec and raw != spec["expected_formula"]:
            raise ValueError("formula mismatch: " + repr(raw))
        return cached, raw

    def trace(self):
        all_claims = self.m.get("claims", [])
        if "trace_graph" in self.m:
            from trace_graph import GraphAudit, effective_claims
            graph = GraphAudit(self.m, self.root,require_substantive=True,substantive_pages=self.scope).build()
            self.rows.extend({**row, "page": ""} for row in graph.rows)
            all_claims = effective_claims(self.m)
        ids = [c["id"] for c in all_claims]
        self.check("claim_ids", len(ids) == len(set(ids)) and bool(ids), "Claim ID 非空且唯一")
        page_ids = [p["id"] for p in self.m["pages"]]
        extra_units = set((self.m.get("trace_graph") or {}).get("units", {}))
        self.check("claim_page_membership", all(page_ids.count(c["page"]) == 1 or c["page"] in extra_units for c in all_claims), "全部 Claim 均须指向登记页或显式补充单元；范围过滤不能隐藏游离 Claim")
        claims = [c for c in all_claims if c["page"] in self.scope]
        for c in claims:
            pid, cid = c["page"], c["id"]
            page = self.base["pages"].get(pid, {})
            self.check("report_quote", bool(c.get("report_quote")) and normalize(c.get("report_quote", "")) in normalize(page.get("heading", "") + "\n" + page.get("body", "")), cid, pid)
            source_ids = c.get("source_ids", [])
            self.check("source_ids", bool(source_ids), cid, pid)
            for sid in source_ids:
                s = self.m.get("sources", {}).get(sid)
                if not self.check("source_registered", bool(s), cid + " → " + sid, pid): continue
                self.file(s, "source_hash", pid)
                self.check("source_pin", bool(s.get("sha256")), sid + " 原文需固定哈希", pid)
                self.check("source_audit", s.get("review_status") == "verified" and bool(s.get("audit_basis")), sid + ": " + s.get("audit_basis", "未人工核验原文"), pid, unknown=True)
                internal = s.get("citation_policy") == "internal" and s.get("source_type") in {"client_input", "project_definition", "research_process_record"}
                if internal:
                    self.check("internal_source_basis", bool(s.get("audit_basis")) and bool(s.get("source_type")), sid + " 内部定义／输入，按登记依据追溯，不作为公开研究APA", pid)
                else:
                    graph_spec = self.m.get("trace_graph")
                    direct = not graph_spec or any(u.get("claim_id") == cid and u.get("citation_policy") != "audit_only" and u.get("from", "").startswith("Evidence:")
                        and graph_spec["evidence"].get(u["from"].split(":", 1)[1], {}).get("source_id") == sid
                        and graph_spec["locations"].get(u.get("location_id"), {}).get("page") == pid
                        for u in graph_spec.get("uses", []))
                    if direct:
                        listed = False
                        if graph_spec:
                            part=next(p.get("part") for p in self.m["pages"] if p["id"]==pid)
                            part_pages={p["id"] for p in self.m["pages"] if p.get("part")==part}
                            for ref in graph_spec.get("source_references",[]):
                                ref_source=self.m["sources"].get(ref.get("source_id"),{})
                                same=ref.get("source_id")==sid or bool(s.get("same_document_group") and s.get("same_document_group")==ref_source.get("same_document_group"))
                                loc=graph_spec["locations"].get(ref.get("location_id"),{})
                                if same and ref.get("purpose")=="bibliography" and loc.get("page") in part_pages and len(loc.get("quote",""))>40:
                                    listed=True;break
                        self.check("apa_in_page", bool(s.get("apa")) and (listed or normalize(s.get("apa", "")) in normalize(page.get("body", ""))), sid, pid)
                for token in [t["token"] for t in c.get("source_checks", []) if t["source_id"] == sid]:
                    try:
                        src = self.path(s["path"])
                        if s.get("extraction") or src.suffix.lower() not in {".html", ".txt", ".md", ".xml", ".json", ".csv"}:
                            extraction = s.get("extraction", {})
                            if not extraction.get("sha256") or extraction.get("original_sha256") != digest(src):
                                raise ValueError("binary source requires hash-bound audited extraction")
                            if not self.file(extraction, "source_extraction", pid):
                                raise ValueError("extraction missing or changed")
                            src = self.path(extraction["path"])
                        self.check("source_token", normalize(token) in normalize(src.read_text(encoding="utf-8-sig")), sid + " / " + token, pid)
                    except (OSError, UnicodeError, ValueError) as e: self.check("source_token", False, str(e), pid, True)
            self.check("source_locator", bool(c.get("source_checks")) and set(source_ids) == {t.get("source_id") for t in c.get("source_checks", []) if t.get("token")} and all(t.get("source_id") in source_ids and t.get("token") for t in c.get("source_checks", [])), cid + " 每个所列来源均须有非空定位摘录，不得遗漏或使用未登记来源", pid)
            if c.get("kind") in {"direct", "qualitative", "definition", "method", "recommendation"}:
                if c.get("kind") == "direct":
                    self.check("direct_display", bool(c.get("display", {}).get("text")) and normalize(c.get("display", {}).get("text", "")) in normalize(c["report_quote"]) and any(normalize(c.get("display", {}).get("text", "")) in normalize(t["token"]) for t in c.get("source_checks", [])), cid + " 直接引用的展示值应出现于定位摘录", pid)
                continue
            try:
                registered = next((r for p in self.pages if p["id"] == pid for r in p.get("excel", []) if r.get("path") == c["excel"]["path"]), None)
                self.check("excel_registered", bool(registered and registered.get("sha256")), cid + " 底稿应纳入本页交付清单并固定哈希", pid)
                if registered: self.file(registered, "excel_hash", pid)
                value, raw = self.cell(c["excel"])
                self.check("excel_value_formula", True, cid + " " + c["excel"]["sheet"] + "!" + c["excel"]["cell"] + "=" + str(value), pid)
                if isinstance(raw, str) and raw.startswith("="):
                    self.check("output_formula_pin", bool(c["excel"].get("expected_formula")), cid + " 公式输出须固定原公式；缓存不能单独证明计算", pid)
                if c.get("calculation"):
                    calc = c["calculation"]
                    self.check("formula_pinned", bool(c["excel"].get("expected_formula")), cid + " 派生数的原公式必须登记", pid)
                    inputs = {n: self.cell(s)[0] for n, s in calc["inputs"].items()}
                    independent = arithmetic(calc["expression"], inputs)
                    self.check("independent_calculation", math.isclose(independent, value, rel_tol=1e-9, abs_tol=1e-10), cid + " independently=" + str(independent), pid)
                    self.check("formula_present", isinstance(raw, str) and raw.startswith("="), cid + " 派生数须有 Excel 公式", pid)
                d = c["display"]
                rounded = Decimal(str(value * d.get("scale", 1))).quantize(Decimal(10) ** -d["decimals"], rounding=ROUND_HALF_UP)
                supplied = Decimal(re.sub(r"[^0-9.\-]", "", d["text"]))
                self.check("display_value", rounded == supplied and normalize(d["text"]) in normalize(c["report_quote"]), cid + " → " + d["text"], pid)
            except (OSError, KeyError, ValueError, ArithmeticError, ImportError) as e:
                self.check("excel_trace", False, cid + ": " + str(e), pid)
        for page in self.pages:
            registered = [c["id"] for c in claims if c["page"] == page["id"]]
            self.check("page_claim_map", set(page.get("trace_claim_ids", [])) == set(registered), "页面 Claim 清单与实际登记一致", page["id"])
        for books in self.books.values():
            for book in books: book.close()
        self.books.clear()

    def compare(self, candidate):
        if not candidate:
            self.check("candidate", False, "未提供待发布或回读文件", unknown=True)
            return
        other = document(candidate)
        expected_order = [p for p in self.base["order"] if p in self.scope]
        other_order = [p for p in other["order"] if p in self.scope]
        self.check("compare_page_order", expected_order == other_order and (not self.all_scope() or self.base["order"] == other["order"]), "含页序、缺页、额外页")
        if self.all_scope():
            self.check("compare_intro", normalize(self.base["intro"]) == normalize(other["intro"]), "摘要／开场内容差异")
        for pid, bp in self.base["pages"].items():
            if pid not in self.scope: continue
            op = other["pages"].get(pid, {})
            for field in ("heading", "conclusion", "body"):
                a, b = normalize(bp[field]), normalize(op.get(field, ""))
                # Conservative: unclassified body deltas require human review, never auto-pass.
                delta = "一致" if a == b else " | ".join(list(difflib.unified_diff(bp[field].splitlines(), op.get(field, "").splitlines(), n=1))[:30])
                self.check("compare_" + field, bool(a) and a == b, delta, pid)
            nums = lambda s: Counter(re.findall(r"\d+(?:[,.]\d+)*(?:%|B|bn|m|k)?", s))
            self.check("compare_numbers", nums(normalize(bp["body"])) == nums(normalize(op.get("body", ""))), "全文数字标记比较（不是指标语义审计）", pid)
        image_reviewed = self.image_review(candidate)
        self.check("image_inventory", self.base["images"] == other["images"] or image_reviewed, "图片引用列表比较；平台转换可使用绑定双方版本的实际图像审查记录")
        if other["images"] or any(p.get("visuals", {}).get("files") for p in self.pages):
            self.check("external_image_bytes", image_reviewed, "回读含 %d 张图；依赖绑定当前版本的图像字节／渲染审查记录，不由名称自动判定" % len(other["images"]), unknown=True)

    def release(self, candidate):
        self.inventory()
        self.trace()
        self.compare(candidate)
        approved = self.m["baseline"].get("approval", {})
        self.check("text_approval", approved.get("status") == "confirmed" and bool(approved.get("basis")), "批准正文确认依据")
        pub = self.m.get("publication", {})
        auth = pub.get("authorization", {})
        self.check("authorization", auth.get("status") == "granted" and bool(auth.get("basis")) and bool(pub.get("target")) and auth.get("target") == pub.get("target") and bool(candidate) and auth.get("candidate_sha256") == digest(candidate) and auth.get("baseline_sha256") == self.m["baseline"].get("sha256") and set(auth.get("pages", [])) == set(self.scope), "本次授权须绑定目标、批准稿、候选内容及页范围；检查器从不自行写入")
        self.check("nonempty_release", bool(self.pages), "发布范围不得为空")
        for p in self.pages:
            self.check("full_trace_coverage", p.get("trace_coverage") == "complete" and bool(p.get("trace_claim_ids")), "登记样例不等于全页三流审计", p["id"], True)
            page = self.base["pages"].get(p["id"], {})
            page_hash = text_digest(page.get("heading", "") + page.get("body", ""))
            self.check("page_approval", p.get("confirmation", {}).get("status") == "confirmed" and p.get("confirmation", {}).get("text_sha256") == page_hash, "正文确认须绑定本页内容", p["id"])
            review = p.get("semantic_review", {})
            self.check("semantic_review", review.get("status") == "verified" and bool(review.get("basis")) and review.get("text_sha256") == page_hash, "证据适配／因果／全量 Claim 覆盖需绑定当前正文的人工审计记录", p["id"], True)
            visual = p.get("visuals", {})
            self.check("visual_approval", visual.get("status") in {"confirmed", "not_applicable"} and bool(visual.get("basis")), "图表实际版本需用户 Review", p["id"], True)

    def save(self, output, mode, candidate=None):
        out = Path(output)
        out.mkdir(parents=True, exist_ok=True)
        issues = [r for r in self.rows if r["status"] != "PASS"]
        result = {"checked_at": datetime.now(timezone.utc).isoformat(), "mode": mode,
                  "manifest_sha256": digest(self.manifest_path), "baseline_sha256": digest(self.path(self.m["baseline"]["path"])) if self.m["baseline"].get("path") else None,
                  "candidate_sha256": digest(candidate) if candidate else None,
                  "pages": self.scope,
                  "result": "PASS" if not issues else "NEEDS_REVIEW", "checks": self.rows,
                  "source_review_status": "records_complete" if any(r["check"].startswith("substance_") for r in self.rows) and not any(r["status"]!="PASS" for r in self.rows if r["check"].startswith("substance_")) else ("needs_review" if "trace_graph" in self.m else "legacy_checks_only"),
                  "limits": "图谱模式强制核对原文逐字段审查记录；agent仍须实际理解原文、核对口径和推导。不重算或修改原工作簿，不发布。"}
        (out / (mode + ".json")).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        lines = ["# " + mode + " 检查", "", "结果：" + result["result"], "", result["limits"], "", "| 页 | 检查 | 状态 | 详情 |", "|---|---|---|---|"]
        for r in self.rows:
            detail = r["detail"].replace("|", "／").replace("\n", " ")
            lines.append("| %s | %s | %s | %s |" % (r["page"], r["check"], r["status"], detail))
        (out / (mode + ".md")).write_text("\n".join(lines) + "\n", encoding="utf-8")
        return 0 if not issues else 1


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--mode", choices=["inventory", "trace", "compare", "release"], required=True)
    parser.add_argument("--candidate")
    parser.add_argument("--pages", help="Comma-separated page IDs; omit for full manifest")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    audit = Audit(args.manifest, args.pages.split(",") if args.pages else None)
    try:
        getattr(audit, args.mode)(args.candidate) if args.mode in {"compare", "release"} else getattr(audit, args.mode)()
    except Exception as e:
        audit.check("execution_error", False, type(e).__name__ + ": " + str(e))
    code = audit.save(args.out, args.mode, args.candidate)
    print(json.dumps({"mode": args.mode, "result": "PASS" if code == 0 else "NEEDS_REVIEW", "checks": len(audit.rows), "issues": sum(r["status"] != "PASS" for r in audit.rows)}, ensure_ascii=False))
    return code


if __name__ == "__main__":
    sys.exit(main())
